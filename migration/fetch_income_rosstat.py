"""
migration/fetch_income_rosstat.py

Загружает показатель 1.2 «Среднедушевые денежные доходы населения» из Excel Росстата.

Источник: https://rosstat.gov.ru/folder/13397
Файл:     urov_10kv_{N}kv-{YYYY}.xlsx  (обновляется ежеквартально)
Лист:     СДД_РФ

Структура файла:
  YYYY год         ← маркер года
  1 квартал   <value>  ← period_date = YYYY-01-01
  2 квартал   <value>  ← period_date = YYYY-04-01
  3 квартал   <value>  ← period_date = YYYY-07-01
  4 квартал   <value>  ← period_date = YYYY-10-01
  Год         <value>  ← пропускаем (идёт в 1.2.y, а не в 1.2)
  1 полугодие         ← пропускаем
  9 месяцев           ← пропускаем

Единица: рублей / месяц (столбец B, числовое значение).

Запуск:
  source venv/bin/activate
  python migration/fetch_income_rosstat.py [--dry-run]

Использует: requests, openpyxl, psycopg2, python-dotenv
ЗАПУСК ТОЛЬКО ЛОКАЛЬНО — rosstat.gov.ru может быть недоступен с зарубежных IP (SSL).
"""

import argparse
import io
import logging
import os
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
import psycopg2
import psycopg2.extras
import requests
from dotenv import load_dotenv

for _env in [Path(__file__).parent / ".env", Path(__file__).parent.parent / ".env"]:
    if _env.exists():
        load_dotenv(_env)
        break

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

PAGE_URL = "https://rosstat.gov.ru/folder/13397"
# Паттерн ссылки на файл среднедушевых доходов по РФ
FILE_PATTERN = re.compile(r"/storage/mediabank/(urov_10kv_\d+kv-\d+\.xlsx)", re.I)

QUARTER_MAP = {
    "1 квартал": (1,  "Q1"),
    "2 квартал": (4,  "Q2"),
    "3 квартал": (7,  "Q3"),
    "4 квартал": (10, "Q4"),
}
SKIP_LABELS = {"Год", "1 полугодие", "9 месяцев"}

INDICATOR_CODE = "1.2"
INDICATOR_NAME = "Среднедушевые денежные доходы населения"


# ── HTTP ────────────────────────────────────────────────────────────────────

def _session() -> requests.Session:
    s = requests.Session()
    s.verify = False          # rosstat.gov.ru имеет проблемы с SSL-цепочкой
    s.headers.update({
        "User-Agent": "Mozilla/5.0 (compatible; realestate-dashboard/1.0)",
    })
    return s


def find_file_url(session: requests.Session) -> str:
    """Находит актуальный URL xlsx-файла на странице /folder/13397."""
    log.info(f"Запрашиваем страницу {PAGE_URL} ...")
    r = session.get(PAGE_URL, timeout=30)
    r.raise_for_status()

    matches = FILE_PATTERN.findall(r.text)
    if not matches:
        raise RuntimeError(
            f"Не нашли ссылку urov_10kv_*.xlsx на странице {PAGE_URL}. "
            "Возможно, Росстат изменил имя файла."
        )
    # Берём первое вхождение — оно самое свежее
    filename = matches[0]
    url = f"https://rosstat.gov.ru/storage/mediabank/{filename}"
    log.info(f"Найден файл: {filename}")
    return url


def download_xlsx(session: requests.Session, url: str) -> bytes:
    log.info(f"Скачиваем {url} ...")
    r = session.get(url, timeout=60)
    r.raise_for_status()
    log.info(f"Скачано {len(r.content):,} байт")
    return r.content


# ── Парсинг Excel ────────────────────────────────────────────────────────────

def parse_xlsx(content: bytes) -> list[dict]:
    """
    Возвращает список записей: {year, quarter, month, period_label, value}.
    Парсит только квартальные строки, «Год» и промежуточные итоги пропускает.
    """
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

    # Ищем лист данных (не «Содержание»)
    ws = None
    for name in wb.sheetnames:
        if name.strip() != "Содержание":
            ws = wb[name]
            break
    if ws is None:
        raise RuntimeError("Не нашли лист данных в xlsx.")
    log.info(f"Парсим лист «{ws.title}»")

    records = []
    current_year: int | None = None

    for row in ws.iter_rows(values_only=True):
        label = str(row[0]).strip() if row[0] is not None else ""
        raw_value = row[1] if len(row) > 1 else None

        # Маркер года: «2024 год», «2025 год 3)» и т.п.
        m = re.match(r"^(\d{4})\s*год", label)
        if m:
            current_year = int(m.group(1))
            continue

        if current_year is None:
            continue

        # Пропускаем строки, которые не нужны
        if label in SKIP_LABELS or not label:
            continue

        q_info = QUARTER_MAP.get(label)
        if q_info is None:
            # Неизвестная метка — пропускаем молча
            continue

        start_month, quarter_label = q_info

        # Значение: пробуем привести к float
        if raw_value is None:
            continue
        try:
            value = float(str(raw_value).replace(",", ".").split()[0])
        except (ValueError, IndexError):
            log.debug(f"Пропускаем нечисловое значение: {label} {current_year} → {raw_value!r}")
            continue

        records.append({
            "year":         current_year,
            "quarter":      quarter_label,
            "month":        start_month,
            "period_label": f"{quarter_label} {current_year}",
            "value":        round(value, 4),
        })

    log.info(f"Распознано квартальных записей: {len(records)}")
    return records


# ── БД ────────────────────────────────────────────────────────────────────────

def get_conn():
    return psycopg2.connect(
        host=os.environ.get("DB_HOST", "localhost"),
        port=int(os.environ.get("DB_PORT", 5432)),
        dbname=os.environ.get("DB_NAME", "realestate"),
        user=os.environ.get("DB_USER", "postgres"),
        password=os.environ.get("DB_PASSWORD", ""),
    )


def get_indicator_id(cur, code: str) -> int:
    cur.execute("SELECT id FROM indicators WHERE code = %s", (code,))
    row = cur.fetchone()
    if not row:
        raise ValueError(f"Индикатор '{code}' не найден в БД.")
    return row[0]


def upsert(conn, indicator_id: int, records: list[dict]) -> int:
    rows = [
        (indicator_id, date(r["year"], r["month"], 1), r["period_label"], r["value"], False)
        for r in records
    ]
    with conn.cursor() as cur:
        psycopg2.extras.execute_values(
            cur,
            """
            INSERT INTO data_points
                (indicator_id, period_date, period_label, value, is_preliminary)
            VALUES %s
            ON CONFLICT (indicator_id, period_date) DO UPDATE
                SET value          = EXCLUDED.value,
                    period_label   = EXCLUDED.period_label,
                    is_preliminary = EXCLUDED.is_preliminary
            """,
            rows,
            template="(%s, %s, %s, %s, %s)",
        )
        cur.execute("REFRESH MATERIALIZED VIEW CONCURRENTLY data_points_with_dynamics")
    conn.commit()
    return len(rows)


# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> int:
    parser = argparse.ArgumentParser(
        description=f"Загрузка {INDICATOR_CODE} из Excel Росстата (rosstat.gov.ru/folder/13397)"
    )
    parser.add_argument("--dry-run", action="store_true",
                        help="Не записывать в БД, только показать что нашли")
    args = parser.parse_args()

    session = _session()

    try:
        url = find_file_url(session)
    except Exception as e:
        log.error(f"Ошибка поиска файла: {e}")
        sys.exit(1)

    try:
        content = download_xlsx(session, url)
    except Exception as e:
        log.error(f"Ошибка загрузки файла: {e}")
        sys.exit(1)

    records = parse_xlsx(content)
    if not records:
        log.error("Не найдено ни одной квартальной записи — проверьте структуру файла.")
        sys.exit(1)

    if args.dry_run:
        log.info("=== DRY RUN — первые 10 записей ===")
        for r in records[:10]:
            log.info(f"  {r['period_label']:12s}  {r['value']:>10.2f} руб./мес.")
        log.info(f"Всего: {len(records)} квартальных записей, БД не тронута.")
        return len(records)

    conn = get_conn()
    try:
        with conn.cursor() as cur:
            ind_id = get_indicator_id(cur, INDICATOR_CODE)
        n = upsert(conn, ind_id, records)
        log.info(f"✅ Upserted {n} строк для {INDICATOR_CODE}")
    except Exception as e:
        log.error(f"Ошибка записи в БД: {e}")
        conn.rollback()
        sys.exit(1)
    finally:
        conn.close()

    return n


if __name__ == "__main__":
    main()
